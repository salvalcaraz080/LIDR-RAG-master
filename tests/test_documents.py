"""Tests del módulo de extracción de documentos (app/services/documents.py).

Sin LLM, sin Redis, sin red. Las fixtures se generan en código para ser reproducibles.
fpdf2 (dev) genera PDFs; python-docx genera el DOCX.
"""

from io import BytesIO

import pytest
from docx import Document
from fpdf import FPDF

from app.services.documents import (
    SUPPORTED_EXTENSIONS,
    DocumentExtractionError,
    extract_text,
)

# ---------------------------------------------------------------------------
# Helpers para generar fixtures en memoria
# ---------------------------------------------------------------------------

PDF_SENTINEL_PAGE_1 = "ESTIMATOR_PDF_SENTINEL_PAGINA_1"
PDF_SENTINEL_PAGE_2 = "ESTIMATOR_PDF_SENTINEL_PAGINA_2"
DOCX_SENTINEL = "ESTIMATOR_DOCX_SENTINEL en el documento Word"


def _make_pdf_two_pages() -> bytes:
    """Genera un PDF de 2 páginas con una frase centinela en cada una."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=PDF_SENTINEL_PAGE_1)
    pdf.add_page()
    pdf.cell(text=PDF_SENTINEL_PAGE_2)
    return pdf.output()


def _make_pdf_blank() -> bytes:
    """Genera un PDF de 1 página sin texto (página en blanco)."""
    pdf = FPDF()
    pdf.add_page()
    return pdf.output()


def _make_docx() -> bytes:
    """Genera un DOCX con una frase centinela."""
    doc = Document()
    doc.add_paragraph(DOCX_SENTINEL)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Fixtures pytest
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def pdf_two_pages() -> bytes:
    return _make_pdf_two_pages()


@pytest.fixture(scope="module")
def pdf_blank() -> bytes:
    return _make_pdf_blank()


@pytest.fixture(scope="module")
def docx_file() -> bytes:
    return _make_docx()


# ---------------------------------------------------------------------------
# Casos de prueba
# ---------------------------------------------------------------------------

class TestPdfExtraction:
    def test_contiene_ambas_centinelas(self, pdf_two_pages: bytes) -> None:
        """El texto extraído debe incluir las frases centinela de ambas páginas."""
        text = extract_text(pdf_two_pages, "doc.pdf")
        assert PDF_SENTINEL_PAGE_1 in text
        assert PDF_SENTINEL_PAGE_2 in text

    def test_contiene_marcadores_de_pagina(self, pdf_two_pages: bytes) -> None:
        """Los marcadores '--- Page N ---' deben estar presentes para las 2 páginas."""
        text = extract_text(pdf_two_pages, "doc.pdf")
        assert "--- Page 1 ---" in text
        assert "--- Page 2 ---" in text

    def test_pdf_vacio_levanta_error(self, pdf_blank: bytes) -> None:
        """Un PDF sin texto debe levantar DocumentExtractionError (guarda de vacío)."""
        with pytest.raises(DocumentExtractionError):
            extract_text(pdf_blank, "blank.pdf")

    def test_bytes_corruptos_levanta_error(self) -> None:
        """Bytes inválidos con extensión .pdf no deben propagar excepciones crudas de pypdf."""
        with pytest.raises(DocumentExtractionError):
            extract_text(b"%PDF-1.4 esto no es un pdf valido", "corrupto.pdf")


class TestDocxExtraction:
    def test_contiene_centinela(self, docx_file: bytes) -> None:
        """El texto extraído del DOCX debe incluir la frase centinela."""
        text = extract_text(docx_file, "doc.docx")
        assert DOCX_SENTINEL in text


class TestExtensionDispatch:
    def test_extension_no_soportada_levanta_error(self) -> None:
        """Una extensión fuera de SUPPORTED_EXTENSIONS debe levantar DocumentExtractionError."""
        with pytest.raises(DocumentExtractionError, match="txt"):
            extract_text(b"contenido cualquiera", "foo.txt")

    def test_extension_case_insensitive(self, pdf_two_pages: bytes) -> None:
        """El dispatch debe ignorar mayúsculas en la extensión."""
        text = extract_text(pdf_two_pages, "DOC.PDF")
        assert PDF_SENTINEL_PAGE_1 in text

    def test_supported_extensions_contiene_pdf_y_docx(self) -> None:
        assert ".pdf" in SUPPORTED_EXTENSIONS
        assert ".docx" in SUPPORTED_EXTENSIONS
